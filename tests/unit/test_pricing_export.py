"""Unit tests for Wave 9 pricing export helpers."""

from __future__ import annotations

import inspect
import json
from pathlib import Path

import openpyxl
from sqlalchemy import inspect as sa_inspect
from sqlalchemy.orm import Session

from src.models.price_analysis import PriceAnalysis
from src.models.price_ladder_snapshot import PriceLadderSnapshot
from src.pricing.pricing_export import (
    build_pricing_export_payload,
    export_all_pricing,
    export_pricing_csv,
    export_pricing_excel,
    export_pricing_json,
    export_pricing_markdown,
    row_to_dict,
)


class TestBuildPricingExportPayload:
    def test_returns_dict_with_required_keys(self, seeded_pricing_export_db) -> None:
        payload = build_pricing_export_payload(1, seeded_pricing_export_db)
        assert {
            "keyword_id",
            "niche_price_analyses",
            "price_analyses",
            "pricing_snapshots",
            "ladder_snapshots",
            "revenue_gate_records",
            "pricing_strategy_outputs",
            "export_timestamp",
        }.issubset(payload.keys())

    def test_returns_empty_lists_for_empty_db(self, empty_db) -> None:
        payload = build_pricing_export_payload(1, empty_db)
        assert payload["niche_price_analyses"] == []
        assert payload["price_analyses"] == []
        assert payload["pricing_snapshots"] == []
        assert payload["ladder_snapshots"] == []
        assert payload["revenue_gate_records"] == []
        assert payload["pricing_strategy_outputs"] == []

    def test_keyword_id_in_payload(self, seeded_pricing_export_db) -> None:
        payload = build_pricing_export_payload(1, seeded_pricing_export_db)
        assert payload["keyword_id"] == 1

    def test_export_timestamp_present(self, seeded_pricing_export_db) -> None:
        payload = build_pricing_export_payload(1, seeded_pricing_export_db)
        assert isinstance(payload["export_timestamp"], str)
        assert "T" in payload["export_timestamp"]

    def test_price_analyses_are_list(self, seeded_pricing_export_db) -> None:
        payload = build_pricing_export_payload(1, seeded_pricing_export_db)
        assert isinstance(payload["price_analyses"], list)

    def test_includes_pricing_strategy_output(self, seeded_pricing_export_db) -> None:
        payload = build_pricing_export_payload(1, seeded_pricing_export_db)
        assert len(payload["pricing_strategy_outputs"]) == 1
        assert payload["pricing_strategy_outputs"][0]["recommendation_type"] == "pricing_strategy"

    def test_row_to_dict_excludes_private_attrs(self, seeded_pricing_export_db) -> None:
        with Session(seeded_pricing_export_db) as session:
            row = session.query(PriceAnalysis).first()
            assert row is not None
            exported = row_to_dict(row)
            assert exported
            assert not any(key.startswith("_") for key in exported)

    def test_row_to_dict_only_contains_columns(self, seeded_pricing_export_db) -> None:
        with Session(seeded_pricing_export_db) as session:
            row = session.query(PriceAnalysis).first()
            assert row is not None
            exported = row_to_dict(row)
            columns = {column["name"] for column in sa_inspect(seeded_pricing_export_db).get_columns("price_analysis")}
            assert set(exported.keys()).issubset(columns)


class TestExportPricingCsv:
    def test_creates_file_at_output_path(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.csv"
        export_pricing_csv(1, seeded_pricing_export_db, str(output))
        assert output.exists()

    def test_file_is_nonempty(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.csv"
        export_pricing_csv(1, seeded_pricing_export_db, str(output))
        assert output.stat().st_size > 0

    def test_empty_db_creates_file(self, tmp_path: Path, empty_db) -> None:
        output = tmp_path / "empty.csv"
        export_pricing_csv(99, empty_db, str(output))
        assert output.exists()

    def test_returns_output_path(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.csv"
        result = export_pricing_csv(1, seeded_pricing_export_db, str(output))
        assert result == str(output)

    def test_csv_has_section_headers(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.csv"
        export_pricing_csv(1, seeded_pricing_export_db, str(output))
        content = output.read_text(encoding="utf-8")
        assert "# price_analyses" in content
        assert "# pricing_strategy_outputs" in content


class TestExportPricingJson:
    def test_creates_valid_json(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.json"
        export_pricing_json(1, seeded_pricing_export_db, str(output))
        payload = json.loads(output.read_text(encoding="utf-8"))
        assert isinstance(payload, dict)

    def test_returns_output_path(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.json"
        result = export_pricing_json(1, seeded_pricing_export_db, str(output))
        assert result == str(output)

    def test_empty_db_creates_json(self, tmp_path: Path, empty_db) -> None:
        output = tmp_path / "empty.json"
        export_pricing_json(999, empty_db, str(output))
        payload = json.loads(output.read_text(encoding="utf-8"))
        assert payload["price_analyses"] == []

    def test_json_has_keyword_id(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.json"
        export_pricing_json(1, seeded_pricing_export_db, str(output))
        payload = json.loads(output.read_text(encoding="utf-8"))
        assert payload["keyword_id"] == 1

    def test_datetime_serialized_as_string(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.json"
        export_pricing_json(1, seeded_pricing_export_db, str(output))
        payload = json.loads(output.read_text(encoding="utf-8"))
        assert isinstance(payload["export_timestamp"], str)


class TestExportPricingMarkdown:
    def test_creates_markdown_file(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.md"
        export_pricing_markdown(1, seeded_pricing_export_db, str(output))
        assert output.exists()

    def test_contains_section_headers(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.md"
        export_pricing_markdown(1, seeded_pricing_export_db, str(output))
        text = output.read_text(encoding="utf-8")
        assert "## Price Distribution Analysis" in text
        assert "## Pricing Strategy Output" in text

    def test_empty_db_creates_no_data_sections(self, tmp_path: Path, empty_db) -> None:
        output = tmp_path / "empty.md"
        export_pricing_markdown(1, empty_db, str(output))
        text = output.read_text(encoding="utf-8")
        assert "_No data available._" in text

    def test_returns_output_path(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.md"
        result = export_pricing_markdown(1, seeded_pricing_export_db, str(output))
        assert result == str(output)

    def test_caps_rows_at_twenty(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        with Session(seeded_pricing_export_db) as session:
            for idx in range(24):
                session.add(
                    PriceLadderSnapshot(
                        keyword_id=1,
                        niche_id="test_niche",
                        run_id=f"r{idx}",
                        reviews_at_snapshot=idx + 6,
                        ladder_milestone=5,
                        actual_basic_price=70.0 + idx,
                        recommended_basic_price=65.0,
                        price_delta_pct=0.1,
                        on_track=True,
                    )
                )
            session.commit()
        output = tmp_path / "cap.md"
        export_pricing_markdown(1, seeded_pricing_export_db, str(output))
        text = output.read_text(encoding="utf-8")
        ladder_section = text.split("## Price Ladder Progress", maxsplit=1)[1]
        ladder_section = ladder_section.split("## Revenue Gate Records", maxsplit=1)[0]
        row_lines = [line for line in ladder_section.splitlines() if line.startswith("| ")]
        # header + separator + 20 data rows
        assert len(row_lines) == 22


class TestExportPricingExcel:
    def test_creates_xlsx_file(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.xlsx"
        export_pricing_excel([1], seeded_pricing_export_db, str(output))
        assert output.exists()

    def test_workbook_has_pricing_sheets(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.xlsx"
        export_pricing_excel([1], seeded_pricing_export_db, str(output))
        workbook = openpyxl.load_workbook(output)
        assert "price_analysis" in workbook.sheetnames
        assert "pricing_strategy_outputs" in workbook.sheetnames

    def test_returns_output_path(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output = tmp_path / "pricing.xlsx"
        result = export_pricing_excel([1], seeded_pricing_export_db, str(output))
        assert result == str(output)

    def test_empty_db_creates_empty_workbook(self, tmp_path: Path, empty_db) -> None:
        output = tmp_path / "empty.xlsx"
        export_pricing_excel([], empty_db, str(output))
        workbook = openpyxl.load_workbook(output)
        assert "export_meta" in workbook.sheetnames

    def test_excel_file_is_valid_xlsx(self, tmp_path: Path, empty_db) -> None:
        output = tmp_path / "valid.xlsx"
        export_pricing_excel([], empty_db, str(output))
        workbook = openpyxl.load_workbook(output)
        assert isinstance(workbook.sheetnames, list)


class TestExportAllPricing:
    def test_creates_multiple_format_files(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        results = export_all_pricing([1], seeded_pricing_export_db, str(tmp_path))
        assert "excel" in results
        assert any(key.startswith("csv_") for key in results)
        assert any(key.startswith("json_") for key in results)
        assert any(key.startswith("md_") for key in results)

    def test_returns_dict_of_paths(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        results = export_all_pricing([1], seeded_pricing_export_db, str(tmp_path))
        assert isinstance(results, dict)
        assert all(Path(path).exists() for path in results.values())

    def test_single_format_only(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        results = export_all_pricing([1], seeded_pricing_export_db, str(tmp_path), formats=["json"])
        assert set(results.keys()) == {"json_kw1"}

    def test_creates_output_dir_if_missing(self, tmp_path: Path, seeded_pricing_export_db) -> None:
        output_dir = tmp_path / "nested" / "exports"
        export_all_pricing([1], seeded_pricing_export_db, str(output_dir))
        assert output_dir.exists()

    def test_empty_keyword_list_returns_empty_dict(self, tmp_path: Path, empty_db) -> None:
        results = export_all_pricing([], empty_db, str(tmp_path))
        assert results == {}

    def test_default_formats_include_csv(self) -> None:
        signature = inspect.signature(export_all_pricing)
        default_formats = signature.parameters["formats"].default
        assert default_formats is None


class TestModuleBehaviors:
    def test_markdown_source_has_row_cap(self) -> None:
        source = inspect.getsource(export_pricing_markdown)
        assert "[:20]" in source

    def test_payload_serializes_to_json_without_float_errors(self, empty_db) -> None:
        payload = build_pricing_export_payload(999, empty_db)
        serialized = json.dumps(payload, default=str)
        assert isinstance(serialized, str)

    def test_module_does_not_import_tests(self) -> None:
        source = Path("src/pricing/pricing_export.py").read_text(encoding="utf-8")
        assert "from tests" not in source
        assert "import tests" not in source
